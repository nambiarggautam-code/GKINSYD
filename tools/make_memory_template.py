#!/usr/bin/env python3
"""Fillable Word memory template (one section per place, grouped by the 9 NFC tag areas).

  python3 tools/make_memory_template.py [out.docx]
Needs python-docx. Fill-in boxes are Word content controls ("click and type").
"""
import json, sys, copy
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "Our-Sydney-Story-Memory-Template.docx"
CREAM, CREAM2, INK, BURG, MUTED = "F6F1E8", "EDE5D7", "242424", "783D46", "6B6560"
SERIF, SANS = "Georgia", "Calibri"

MEMORY_PROMPTS = [
    "What do you picture first when you think of this place? What happened that day?",
    "A small detail you remember: the weather, a smell, something you ate, a sound.",
    "What were you doing here, and why has it stayed with you?",
    "Was it a first visit, or somewhere you kept coming back to? What made it special?",
    "Something funny, sweet or unexpected that happened here.",
]
LOVE_PROMPTS = [
    "Something she said or did here that you noticed.",
    "What this day showed you about her.",
    "A quality of hers you saw clearly here: kindness, humour, patience, courage...",
    "How she made this place feel different for you.",
]
ONEDAY = "Optional. Something you would like to do here, or somewhere this place makes you want to go next."
LEAVEOUT = "Optional. Anything that should NOT go on the website (it is public): names, faces, addresses, private details."


def rgb(h):
    return RGBColor.from_string(h)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def borders(table, color="C9BFAF", size=4):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{e}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0"); el.set(qn("w:color"), color); b.append(el)
    tblPr.append(b)


def row_height(row, cm, exact=False):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight"); h.set(qn("w:val"), str(int(cm * 567))); h.set(qn("w:hRule"), "exact" if exact else "atLeast")
    trPr.append(h)
    cs = OxmlElement("w:cantSplit"); trPr.append(cs)


def run(p, text, size=10.5, bold=False, italic=False, color=INK, font=SANS):
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
    r.font.color.rgb = rgb(color); r.font.name = font
    r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    return r


def para(container, text="", size=10.5, bold=False, italic=False, color=INK, font=SANS, before=0, after=4,
         align=None, keep=False):
    p = container.add_paragraph()
    pf = p.paragraph_format; pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = 1.15
    if keep:
        pf.keep_with_next = True
    if align:
        p.alignment = align
    if text:
        run(p, text, size, bold, italic, color, font)
    return p


_sdt_id = [1000]


def fill_box(cell, placeholder, tag):
    """Replace the cell's paragraph with a block-level rich-text content control showing grey placeholder text."""
    tc = cell._tc
    for p in list(tc.findall(qn("w:p"))):
        tc.remove(p)
    _sdt_id[0] += 1
    sdt = OxmlElement("w:sdt")
    pr = OxmlElement("w:sdtPr")
    alias = OxmlElement("w:alias"); alias.set(qn("w:val"), tag[:60]); pr.append(alias)
    t = OxmlElement("w:tag"); t.set(qn("w:val"), tag[:60]); pr.append(t)
    i = OxmlElement("w:id"); i.set(qn("w:val"), str(_sdt_id[0])); pr.append(i)
    pr.append(OxmlElement("w:showingPlcHdr"))
    sdt.append(pr)
    content = OxmlElement("w:sdtContent")
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr"); sp = OxmlElement("w:spacing"); sp.set(qn("w:before"), "40"); sp.set(qn("w:after"), "40")
    ppr.append(sp); p.append(ppr)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    f = OxmlElement("w:rFonts"); f.set(qn("w:ascii"), SANS); f.set(qn("w:hAnsi"), SANS); rpr.append(f)
    rpr.append(OxmlElement("w:i"))
    c = OxmlElement("w:color"); c.set(qn("w:val"), "9A928A"); rpr.append(c)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "19"); rpr.append(sz)
    r.append(rpr)
    tx = OxmlElement("w:t"); tx.text = placeholder; tx.set(qn("xml:space"), "preserve"); r.append(tx)
    p.append(r); content.append(p); sdt.append(content)
    tc.append(sdt)


def label_cell(cell, title, hint=None):
    shade(cell, CREAM2)
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(2)
    run(p, title, 9.5, bold=True, color=INK)
    if hint:
        q = cell.add_paragraph(); q.paragraph_format.space_after = Pt(2)
        run(q, hint, 8, italic=True, color=MUTED)


def field_table(doc, rows, tagbase):
    """rows: (label, hint, placeholder, min_height_cm)"""
    tbl = doc.add_table(rows=len(rows), cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    borders(tbl)
    for r, (lab, hint, ph, h) in zip(tbl.rows, rows):
        r.cells[0].width = Cm(4.2); r.cells[1].width = Cm(12.8)
        label_cell(r.cells[0], lab, hint)
        fill_box(r.cells[1], ph, f"{tagbase} {lab}")
        row_height(r, h)
    return tbl


def heading(doc, text, size, color=INK, before=0, after=6, italic=False, keep=True, align=None):
    p = para(doc, text, size, italic=italic, color=color, font=SERIF, before=before, after=after, keep=keep, align=align)
    return p


def rule(doc):
    p = para(doc, "", after=6)
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "4"); b.set(qn("w:space"), "1"); b.set(qn("w:color"), "C9BFAF")
    bdr.append(b); pPr.append(bdr)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def main():
    data = json.load(open(ROOT / "private" / "locations.json" if (ROOT / "private" / "locations.json").exists() else ROOT / "data" / "locations.json"))
    locs = {l["slug"]: l for l in data["locations"]}
    cl = {c["slug"]: c for c in data["clusters"]}
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = Cm(1.8); sec.bottom_margin = Cm(1.8)
    # cream page colour (shown on screen and in the PDF)
    bg = OxmlElement("w:background"); bg.set(qn("w:color"), CREAM)
    doc.element.insert(0, bg)
    st = doc.settings.element
    dbs = OxmlElement("w:displayBackgroundShape"); st.insert(0, dbs)
    normal = doc.styles["Normal"]; normal.font.name = SANS; normal.font.size = Pt(10.5); normal.font.color.rgb = rgb(INK)
    # footer
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(fp, "Our Sydney Story · memory template", 8, italic=True, color=MUTED, font=SERIF)

    # ---------------- page 1: how to fill it in
    heading(doc, "Our Sydney Story", 30, before=10, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    heading(doc, "A memory for every place on the map", 14, BURG, after=10, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    rule(doc)
    para(doc, "This document holds everything the website needs. Fill in as much or as little as you like: "
              "anything left empty can be added later, and nothing has to be perfect. Write the way you talk. "
              "Small, specific things usually say more than big words.", 10.5, after=8)
    heading(doc, "How to fill it in", 15, before=4)
    steps = [
        ("Pick one photo for each place.", "Any photo you love from there. Name the file with the place number, "
         "for example 07.jpg or 07-taronga-zoo.jpg (each section suggests a name). Phone photos are fine."),
        ("Write the memory: 2 to 4 sentences.", "What happened, what you noticed, why it stayed with you."),
        ("Write \u201cWhat I love about you\u201d: 2 to 3 sentences.", "Written to her. What this place or day showed you about her."),
        ("Optional: \u201cOne day, we\u2019ll\u2026\u201d", "One line about something to do together in the future. Leave it empty if nothing comes to mind."),
        ("Optional: the date.", "A rough date is fine (\u201cMarch 2026\u201d). Leave it empty if you are not sure."),
        ("Note anything to leave out.", "The website is public: anyone with the link can see it. Mention anything "
         "that should stay private."),
        ("Fill in the last pages.", "The letter (\u201cThe Best Part of This Journey\u201d), her name or nickname, "
         "how you sign off, the homepage introduction and a one-line intro for each area."),
        ("Send it back.", "Send this document and the photos (named by number) back in our chat, and the site gets updated."),
    ]
    for i, (a, b) in enumerate(steps, 1):
        p = para(doc, "", after=4)
        p.paragraph_format.left_indent = Cm(0.8); p.paragraph_format.first_line_indent = Cm(-0.8)
        run(p, f"{i}.\u2002", 10.5, bold=True, color=BURG, font=SERIF)
        run(p, a + " ", 10.5, bold=True)
        run(p, b, 10.5)
    heading(doc, "If you get stuck", 13, BURG, before=8, italic=True)
    for t in ["Close your eyes and picture the place. What is the first thing you see?",
              "Think of one moment, not the whole day.",
              "What did she say, do or notice that you remember?",
              "How did you feel walking away from there?",
              "It is fine to skip a place, or to write just one sentence."]:
        p = para(doc, "", after=2); p.paragraph_format.left_indent = Cm(0.8)
        run(p, "\u2013\u2002" + t, 10, italic=True, color=INK)
    heading(doc, "How the map works", 13, BURG, before=8, italic=True)
    para(doc, "The framed map has 9 hidden NFC tags. Tapping a phone on one opens a page for that area of the map, "
              "which lists the places in it. Each place below shows its area, its number and its web page name "
              "(for example memories/taronga-zoo.html). Please don\u2019t rename places; the page names are permanent.",
         10, after=4)

    # ---------------- places by tag area
    city = cl["sydney-city"]["members"]
    n_done = []

    def place(slug, idx_in_page):
        l = locs[slug]; n = l["number"]
        p = heading(doc, f"{n:02d} \u00b7 {l['name']}", 14, BURG, before=8 if idx_in_page else 0, after=0)
        q = para(doc, "", after=4, keep=True)
        run(q, f"Site page: memories/{slug}.html  \u00b7  suggested photo name: {n:02d}-{slug}.jpg", 8, color=MUTED)
        k = len(n_done)
        field_table(doc, [
            ("Date visited", "optional", "e.g. 14 February 2026, or just \u201cautumn 2026\u201d", 0.8),
            ("Photo file name", None, f"e.g. {n:02d}-{slug}.jpg", 0.8),
            ("The memory", "2\u20134 sentences", MEMORY_PROMPTS[k % len(MEMORY_PROMPTS)], 3.0),
            ("What I love about you", "2\u20133 sentences", LOVE_PROMPTS[k % len(LOVE_PROMPTS)], 2.2),
            ("One day, we\u2019ll\u2026", "optional", ONEDAY, 1.0),
            ("Anything to leave out", "optional", LEAVEOUT, 1.0),
        ], f"{n:02d}")
        n_done.append(slug)

    for t in sorted(data["tags"], key=lambda t: t["number"]):
        page_break(doc)
        heading(doc, f"Tag {t['number']} \u00b7 {t['name']}", 22, after=2)
        if t["target"] in cl:
            members = sorted(cl[t["target"]]["members"], key=lambda m: locs[m]["number"])
            para(doc, f"Tapping this tag opens memories/{t['target']}.html, which lists {len(members)} places.",
                 9.5, italic=True, color=MUTED, after=2)
        else:
            members = [t["target"]]
            para(doc, f"Tapping this tag opens the place\u2019s own page, memories/{t['target']}.html.",
                 9.5, italic=True, color=MUTED, after=2)
        rule(doc)
        if t["target"] == "city-and-east":
            heading(doc, "In the city (one dot on the map)", 13, INK, before=2, italic=True)
            for i, s in enumerate(sorted(city, key=lambda m: locs[m]["number"])):
                if i and i % 2 == 0:
                    page_break(doc)
                place(s, i % 2)
            page_break(doc)
            heading(doc, "Around the harbour and the east", 13, INK, before=2, italic=True)
            rest = [m for m in members if m not in city]
            for i, s in enumerate(rest):
                if i and i % 2 == 0:
                    page_break(doc)
                place(s, i % 2)
        else:
            for i, s in enumerate(members):
                if i and i % 2 == 0:
                    page_break(doc)
                place(s, i % 2)

    off = [l["slug"] for l in data["locations"] if l.get("on_map") is False]
    if off:
        page_break(doc)
        heading(doc, "Not on the map", 22, after=2)
        para(doc, "These places have their own page on the website, but no dot or tag on the framed map.",
             9.5, italic=True, color=MUTED, after=2)
        rule(doc)
        for i, s in enumerate(off):
            place(s, i % 2)

    # ---------------- letter
    page_break(doc)
    heading(doc, "The Best Part of This Journey", 24, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    heading(doc, "The letter at the end of the website", 12, BURG, italic=True, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "A few prompts, only if they help. Three or four short paragraphs is plenty.", 10, after=4)
    for t in ["Where were you when you decided to come to Sydney, and how did it feel?",
              "A moment from these months that you think about often.",
              "Something you have noticed about her that you may never have said out loud.",
              "What you are grateful for.",
              "What you are looking forward to."]:
        p = para(doc, "", after=2); p.paragraph_format.left_indent = Cm(0.8)
        run(p, "\u2013\u2002" + t, 10, italic=True)
    para(doc, "", after=4)
    field_table(doc, [
        ("Opening line", "e.g. Dear \u2026,", "Dear ...", 0.9),
        ("The letter", "write freely", "Start wherever feels natural.", 13.0),
        ("Closing line", "e.g. With love,", "With love,", 0.9),
    ], "Letter")

    # ---------------- homepage + names
    page_break(doc)
    heading(doc, "The homepage and a few names", 22, after=2)
    rule(doc)
    field_table(doc, [
        ("Her name or nickname", "as it should appear", "The name you call her.", 0.9),
        ("How you sign off", "your name", "e.g. your first name, or a name only she uses for you", 0.9),
        ("Homepage introduction", "2\u20134 sentences", "A short welcome she reads first: when you arrived in Sydney, "
         "how it felt, and that every place on the map holds a memory of her.", 4.0),
        ("Line under the title", "optional", "Currently \u201cSince February 2026\u201d. Change it or leave as is.", 0.9),
        ("Footer line", "optional", "Currently \u201cEverywhere we\u2019ve been, a little piece of us.\u201d", 0.9),
    ], "Home")

    heading(doc, "One line for each area page", 15, before=14)
    para(doc, "Shown at the top of the page each tag opens. One sentence each, or leave empty. "
              "(Tags 8 and 9 open the place\u2019s own page, so their line is optional.)", 9.5, italic=True, color=MUTED, after=4)
    rows = []
    for t in sorted(data["tags"], key=lambda t: t["number"]):
        rows.append((f"{t['number']}. {t['name']}", None, "One sentence about this part of the map.", 1.0))
    rows.append(("Sydney City", "the city dot\u2019s own page", "Optional. One sentence about the city.", 1.0))
    field_table(doc, rows, "Area")
    para(doc, "", after=6)
    para(doc, "Thank you. Send this file and the photos back in our chat whenever you are ready, even if it is only half done.",
         10, italic=True, color=BURG, font=SERIF, align=WD_ALIGN_PARAGRAPH.CENTER, before=8)
    doc.save(OUT)
    print("wrote", OUT, "places:", len(n_done))


if __name__ == "__main__":
    main()
